"""
document_processor.py
=====================
EvidencePro — Document Upload and Processing Layer
----------------------------------------------------

PURPOSE
-------
Accepts an uploaded forensic document (PDF, image, DOCX) and converts it
to structured text that can be fed into the AI extraction pipeline.

DOCLING INTEGRATION
-------------------
This module uses the open-source Docling library where available.
Docling handles PDF conversion, image OCR, and DOCX extraction in a
unified pipeline.

FALLBACK CHAIN
--------------
1. Docling (full OCR + document understanding) — if installed
2. PyMuPDF / pdfplumber for PDF text extraction — if installed
3. python-docx for DOCX — if installed
4. PIL / pytesseract for images — if installed
5. Plaintext fallback — return empty text, show clear error message

The application NEVER crashes if a dependency is missing.
Each fallback is clearly labelled so the investigator knows what was used.

IMPORTANT NOTICES
-----------------
- Extracted text is ALWAYS presented for investigator review.
- OCR/extraction is imperfect; results must be verified.
- Document upload is an optional convenience feature; the manual workflow
  always remains available.

Usage
-----
    from document_processor import process_uploaded_document, DoclingAvailable

    result = process_uploaded_document(file_bytes, filename)
    # result.text        : extracted text
    # result.method      : how it was extracted ("docling", "pymupdf", etc.)
    # result.success     : True if any text was extracted
    # result.pages       : number of pages detected (if available)
    # result.warning     : human-readable warning if fallback was used
"""

from __future__ import annotations

import io
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


# ---------------------------------------------------------------------------
# DEPENDENCY AVAILABILITY FLAGS
# ---------------------------------------------------------------------------

def _check_docling() -> bool:
    """Return True if docling is installed and importable."""
    try:
        import importlib.util
        return importlib.util.find_spec("docling") is not None
    except Exception:
        return False


def _check_pymupdf() -> bool:
    try:
        import importlib.util
        return (
            importlib.util.find_spec("fitz") is not None or
            importlib.util.find_spec("pymupdf") is not None
        )
    except Exception:
        return False


def _check_pdfplumber() -> bool:
    try:
        import importlib.util
        return importlib.util.find_spec("pdfplumber") is not None
    except Exception:
        return False


def _check_docx() -> bool:
    try:
        import importlib.util
        return importlib.util.find_spec("docx") is not None
    except Exception:
        return False


def _check_pil() -> bool:
    try:
        import importlib.util
        return importlib.util.find_spec("PIL") is not None
    except Exception:
        return False


def _check_pytesseract() -> bool:
    try:
        import importlib.util
        return importlib.util.find_spec("pytesseract") is not None
    except Exception:
        return False


DOCLING_AVAILABLE   = _check_docling()
PYMUPDF_AVAILABLE   = _check_pymupdf()
PDFPLUMBER_AVAILABLE = _check_pdfplumber()
DOCX_AVAILABLE      = _check_docx()
PIL_AVAILABLE       = _check_pil()
TESSERACT_AVAILABLE = _check_pytesseract()

# Supported file extensions
SUPPORTED_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".docx", ".txt"}


# ---------------------------------------------------------------------------
# RESULT DATACLASS
# ---------------------------------------------------------------------------

@dataclass
class DocumentResult:
    """
    Result of processing one uploaded document.

    Attributes
    ----------
    text     : Extracted text content. May be empty string if extraction failed.
    method   : Processing method used (e.g., "docling", "pymupdf", "manual").
    success  : True if meaningful text was extracted.
    pages    : Number of pages detected (0 if unknown).
    warning  : Human-readable warning or status message for the investigator.
    filename : Original filename of the uploaded document.
    """
    text:     str  = ""
    method:   str  = "none"
    success:  bool = False
    pages:    int  = 0
    warning:  str  = ""
    filename: str  = ""


# ---------------------------------------------------------------------------
# PUBLIC API
# ---------------------------------------------------------------------------

def processing_capability_summary() -> str:
    """
    Return a human-readable summary of available document processing methods.
    Used in the UI to inform the investigator what document types are supported.
    """
    if DOCLING_AVAILABLE:
        return "Docling (full OCR + PDF/image/DOCX conversion)"
    elif PYMUPDF_AVAILABLE:
        return "PyMuPDF (PDF text extraction, no OCR)"
    elif PDFPLUMBER_AVAILABLE:
        return "pdfplumber (PDF text extraction, no OCR)"
    elif DOCX_AVAILABLE:
        return "python-docx (DOCX only)"
    elif PIL_AVAILABLE and TESSERACT_AVAILABLE:
        return "pytesseract (image OCR only)"
    else:
        return "No document processing libraries available — manual entry only"


def is_document_processing_available() -> bool:
    """Return True if ANY document processing capability is available."""
    return (
        DOCLING_AVAILABLE
        or PYMUPDF_AVAILABLE
        or PDFPLUMBER_AVAILABLE
        or DOCX_AVAILABLE
        or (PIL_AVAILABLE and TESSERACT_AVAILABLE)
    )


def process_uploaded_document(
    file_bytes: bytes,
    filename:   str,
) -> DocumentResult:
    """
    Process an uploaded document and extract its text content.

    Tries extraction methods in order of preference:
    1. Docling (best — handles PDF, images with OCR, DOCX)
    2. PyMuPDF / pdfplumber for PDFs
    3. python-docx for DOCX
    4. PIL + pytesseract for images
    5. Raw UTF-8 decode for .txt files

    Parameters
    ----------
    file_bytes : Raw bytes of the uploaded file.
    filename   : Original filename (used to determine file type).

    Returns
    -------
    DocumentResult — always returns a result, never raises.
    """
    if not file_bytes:
        return DocumentResult(
            filename=filename,
            warning="Empty file received. Please upload a valid document.",
        )

    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        return DocumentResult(
            filename=filename,
            warning=(
                f"File type '{suffix}' is not supported. "
                f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            ),
        )

    # .txt — direct decode, no library needed
    if suffix == ".txt":
        return _process_txt(file_bytes, filename)

    # Try Docling first (handles all supported types)
    if DOCLING_AVAILABLE:
        result = _process_via_docling(file_bytes, filename, suffix)
        if result.success:
            return result
        # Docling failed — fall through to alternative

    # PDF alternatives
    if suffix == ".pdf":
        if PYMUPDF_AVAILABLE:
            result = _process_pdf_pymupdf(file_bytes, filename)
            if result.success:
                return result
        if PDFPLUMBER_AVAILABLE:
            result = _process_pdf_pdfplumber(file_bytes, filename)
            if result.success:
                return result

    # DOCX alternative
    if suffix == ".docx" and DOCX_AVAILABLE:
        result = _process_docx(file_bytes, filename)
        if result.success:
            return result

    # Image alternatives (OCR)
    if suffix in {".jpg", ".jpeg", ".png"}:
        if PIL_AVAILABLE and TESSERACT_AVAILABLE:
            result = _process_image_tesseract(file_bytes, filename)
            if result.success:
                return result
        elif PIL_AVAILABLE:
            return DocumentResult(
                filename=filename,
                method="pil_no_ocr",
                warning=(
                    "Image uploaded but pytesseract (OCR) is not installed. "
                    "Install pytesseract and Tesseract OCR to extract text from images. "
                    "Please enter the case details manually."
                ),
            )

    # Nothing worked
    missing = _missing_library_message(suffix)
    return DocumentResult(
        filename=filename,
        method="none",
        warning=(
            f"No text could be extracted from '{filename}'. {missing} "
            "Please enter the case details manually using the form below."
        ),
    )


# ---------------------------------------------------------------------------
# PROCESSING METHODS
# ---------------------------------------------------------------------------

def _process_txt(file_bytes: bytes, filename: str) -> DocumentResult:
    """Extract text from a plain-text file."""
    try:
        text = file_bytes.decode("utf-8", errors="replace").strip()
        return DocumentResult(
            text=text,
            method="plaintext",
            success=bool(text),
            pages=1,
            filename=filename,
            warning="" if text else "The text file appears to be empty.",
        )
    except Exception as e:
        return DocumentResult(
            filename=filename,
            warning=f"Could not read text file: {e}",
        )


def _process_via_docling(
    file_bytes: bytes,
    filename:   str,
    suffix:     str,
) -> DocumentResult:
    """
    Use Docling to convert a document to text.

    Docling handles PDF (with OCR if needed), images, and DOCX.
    Writes a temporary file because Docling requires a file path.
    """
    try:
        from docling.document_converter import DocumentConverter  # type: ignore

        with tempfile.NamedTemporaryFile(
            suffix=suffix, delete=False
        ) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            converter = DocumentConverter()
            result = converter.convert(tmp_path)
            # Export to Markdown text (preserves structure well)
            text = result.document.export_to_markdown()
            pages = _estimate_pages(text)
            return DocumentResult(
                text=text.strip(),
                method="docling",
                success=bool(text.strip()),
                pages=pages,
                filename=filename,
                warning=(
                    "Extracted via Docling (document understanding + OCR where required). "
                    "Review all extracted values carefully — AI extraction is not perfect."
                ) if text.strip() else "Docling could not extract text from this document.",
            )
        finally:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass

    except Exception as e:
        return DocumentResult(
            filename=filename,
            method="docling_failed",
            warning=f"Docling extraction failed ({type(e).__name__}). Trying fallback method.",
        )


def _process_pdf_pymupdf(file_bytes: bytes, filename: str) -> DocumentResult:
    """Extract text from a PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # type: ignore (PyMuPDF)

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages_text = []
        for page in doc:
            pages_text.append(page.get_text())
        doc.close()

        text = "\n\n".join(pages_text).strip()
        return DocumentResult(
            text=text,
            method="pymupdf",
            success=bool(text),
            pages=len(pages_text),
            filename=filename,
            warning=(
                "Extracted via PyMuPDF (PDF text layer only — no OCR). "
                "Scanned or image-based PDFs may produce incomplete text."
            ) if text else "PyMuPDF could not find text in this PDF (scanned/image-based?).",
        )
    except Exception as e:
        return DocumentResult(
            filename=filename,
            method="pymupdf_failed",
            warning=f"PyMuPDF extraction failed: {e}",
        )


def _process_pdf_pdfplumber(file_bytes: bytes, filename: str) -> DocumentResult:
    """Extract text from a PDF using pdfplumber."""
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_text = [p.extract_text() or "" for p in pdf.pages]

        text = "\n\n".join(pages_text).strip()
        return DocumentResult(
            text=text,
            method="pdfplumber",
            success=bool(text),
            pages=len(pages_text),
            filename=filename,
            warning=(
                "Extracted via pdfplumber (PDF text layer only — no OCR). "
                "Scanned or image-based PDFs may produce incomplete text."
            ) if text else "pdfplumber could not find text in this PDF.",
        )
    except Exception as e:
        return DocumentResult(
            filename=filename,
            method="pdfplumber_failed",
            warning=f"pdfplumber extraction failed: {e}",
        )


def _process_docx(file_bytes: bytes, filename: str) -> DocumentResult:
    """Extract text from a DOCX file using python-docx."""
    try:
        import docx  # type: ignore (python-docx)

        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        text = "\n\n".join(paragraphs).strip()
        return DocumentResult(
            text=text,
            method="python-docx",
            success=bool(text),
            pages=1,
            filename=filename,
            warning=(
                "Extracted via python-docx. "
                "Review all extracted values — formatting may affect text quality."
            ) if text else "python-docx could not extract text from this file.",
        )
    except Exception as e:
        return DocumentResult(
            filename=filename,
            method="python-docx-failed",
            warning=f"DOCX extraction failed: {e}",
        )


def _process_image_tesseract(file_bytes: bytes, filename: str) -> DocumentResult:
    """OCR an image using PIL + pytesseract."""
    try:
        from PIL import Image  # type: ignore
        import pytesseract  # type: ignore

        img = Image.open(io.BytesIO(file_bytes))
        text = pytesseract.image_to_string(img).strip()
        return DocumentResult(
            text=text,
            method="tesseract-ocr",
            success=bool(text),
            pages=1,
            filename=filename,
            warning=(
                "Extracted via Tesseract OCR. "
                "OCR accuracy depends on image quality. "
                "Review all extracted values carefully."
            ) if text else "Tesseract OCR could not extract text from this image.",
        )
    except Exception as e:
        return DocumentResult(
            filename=filename,
            method="tesseract_failed",
            warning=f"Tesseract OCR failed: {e}",
        )


# ---------------------------------------------------------------------------
# UTILITIES
# ---------------------------------------------------------------------------

def _estimate_pages(text: str) -> int:
    """Estimate page count from extracted text (rough heuristic)."""
    # ~400 words per page is a rough estimate
    words = len(text.split())
    return max(1, words // 400)


def _missing_library_message(suffix: str) -> str:
    """Return a helpful message about what's missing for a given file type."""
    if suffix == ".pdf":
        if not DOCLING_AVAILABLE and not PYMUPDF_AVAILABLE and not PDFPLUMBER_AVAILABLE:
            return (
                "To process PDF files, install: pip install docling "
                "(or pip install pymupdf or pip install pdfplumber)"
            )
    elif suffix in {".jpg", ".jpeg", ".png"}:
        if not DOCLING_AVAILABLE:
            return (
                "To extract text from images, install: pip install docling "
                "(recommended) or pip install pytesseract and install Tesseract OCR."
            )
    elif suffix == ".docx":
        if not DOCLING_AVAILABLE and not DOCX_AVAILABLE:
            return "To process DOCX files, install: pip install docling or pip install python-docx"
    return ""


def explain_document_processing() -> dict:
    """
    Return a structured explanation of the document processing pipeline
    for display in the UI help/info sections.
    """
    return {
        "docling_available": DOCLING_AVAILABLE,
        "pymupdf_available": PYMUPDF_AVAILABLE,
        "pdfplumber_available": PDFPLUMBER_AVAILABLE,
        "docx_available": DOCX_AVAILABLE,
        "tesseract_available": PIL_AVAILABLE and TESSERACT_AVAILABLE,
        "capability_summary": processing_capability_summary(),
        "supported_types": sorted(SUPPORTED_EXTENSIONS),
        "is_available": is_document_processing_available(),
    }
